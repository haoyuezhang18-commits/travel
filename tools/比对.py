#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
比对两份行程 DOCX，逐行列出时间改动、新增行、删除行。

负责人改过稿之后，做后续文档（文档2、汇总表）前必须先跑一次，
否则会拿旧数据去做——只看表格结构和标题是看不出改动的。

用法：
    python3 tools/比对.py 我生成的.docx 负责人改过的.docx
"""
import sys, os, argparse
from docx import Document


def grab(path):
    d = Document(os.path.expanduser(path))
    titles = [p.text.strip() for p in d.paragraphs
              if p.text.strip().startswith('D') and '｜' in p.text]
    out = []
    for i, t in enumerate(d.tables):
        key = titles[i] if i < len(titles) else f'表{i+1}'
        rows = []
        for r in list(t.rows)[1:]:
            cells = [c.text.strip().replace('\n', ' ') for c in r.cells]
            rows.append((cells[0], cells[1] if len(cells) > 1 else ''))
        out.append((key, rows))
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('old', help='我生成的版本')
    ap.add_argument('new', help='负责人改过的版本')
    a = ap.parse_args()

    old, new = grab(a.old), grab(a.new)
    om = {k.split('｜')[0].strip(): v for k, v in old}
    total = 0

    for key, rows in new:
        day = key.split('｜')[0].strip()
        prev = om.get(day)
        if prev is None:
            print(f'\n### {key}　【新增的一天】共 {len(rows)} 行')
            total += 1
            continue
        pn = {n: t for t, n in prev}
        cur = [n for _, n in rows]
        diffs = []
        for t, n in rows:
            if n in pn:
                if pn[n] != t:
                    diffs.append(f'  ⏱ 时间改动　{n[:30]}　{pn[n]} → {t}')
            else:
                diffs.append(f'  ＋ 新增　{t}　{n[:36]}')
        for t, n in prev:
            if n not in cur:
                diffs.append(f'  － 删除　{t}　{n[:36]}')
        if diffs:
            print(f'\n### {key}')
            for x in diffs:
                print(x)
            total += len(diffs)

    print(f'\n共 {total} 处改动。'
          + ('　⚠ 做文档2、汇总表前请按这些改动更新数据。' if total else '　两份一致。'))
    return 0


if __name__ == '__main__':
    sys.exit(main())
