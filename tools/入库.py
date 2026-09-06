#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把交付成品里的某张表，连同图片和排版，搬进模板库对应城市的文件里。

    python3 tools/入库.py 成品.docx --表 3 --目标 "英国旅游通用模板/主题/攀岩馆通用.docx" --标题 "【伦敦】Castle 攀岩"
    python3 tools/入库.py 成品.docx --表 3 --要 "攀岩,冲澡" --目标 ... --标题 ...

目标文件不存在就新建，存在就在末尾追加一张表。图片走 sha1 去重重新登记，
不能直接 relate_to 源 part，否则 partname 撞车产出损坏文件（见 制作指南）。
"""
import argparse, copy, hashlib, io, os, sys, yaml
from docx import Document
from docx.oxml.ns import qn

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
TR, TC, WT = qn('w:tr'), qn('w:tc'), qn('w:t')
HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def template_root(cli=None):
    if cli:
        return os.path.expanduser(cli)
    cfg = os.path.join(HERE, 'config', 'local-paths.yaml')
    if os.path.exists(cfg):
        d = yaml.safe_load(open(cfg)) or {}
        if d.get('template_root'):
            return os.path.expanduser(d['template_root'])
    return os.path.expanduser('~/文件夹/旅游通用模板')


def cell_text(tc):
    return ''.join(t.text or '' for t in tc.iter(WT))


def copy_images(tbl_el, src_part, dst_part, seen):
    """图片二进制交给 get_or_add_image 重新登记，按 sha1 去重分配唯一 partname。
    不能直接 relate_to 源 part：不同模板里都有 image1.png，partname 会撞车。"""
    R_NS = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
    n = 0
    skipped = []
    for el in tbl_el.iter():
        for attr in (R_NS + 'embed', R_NS + 'link', R_NS + 'id'):
            rid = el.get(attr)
            if not rid:
                continue
            try:
                part = src_part.related_parts[rid]
            except KeyError:
                continue
            ct = getattr(part, 'content_type', '') or ''
            if not ct.startswith('image/'):
                continue                      # 超链接、图表等非图片关系跳过
            blob = getattr(part, 'blob', None)
            if not blob:
                continue
            key = hashlib.sha1(blob).hexdigest()
            if key not in seen:
                try:
                    new_rid, _ = dst_part.get_or_add_image(io.BytesIO(blob))
                except Exception:
                    # Word 重存过的文档里可能混进 EMF/WMF 等 python-docx 认不出的格式，
                    # 跳过这一张，别让整次入库失败
                    seen[key] = None
                else:
                    seen[key] = new_rid
            if seen[key] is None:
                skipped.append(ct)
                continue
            el.set(attr, seen[key])
            n += 1
    if skipped:
        print(f'  ⚠ 跳过 {len(skipped)} 张无法解析的图片（{set(skipped)}）')
    return n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('源')
    ap.add_argument('--表', type=int, required=True, help='源文件里第几张表（从 1 数）')
    ap.add_argument('--目标', required=True, help='相对模板库根目录的路径')
    ap.add_argument('--标题', required=True, help='表前面加的标题段，用于打编号索引')
    ap.add_argument('--要', default='', help='只取含这些关键词的行，逗号分隔；留空取整表')
    ap.add_argument('--root', default=None)
    a = ap.parse_args()

    src = Document(os.path.expanduser(a.源))
    tbls = src.tables
    if not 1 <= a.表 <= len(tbls):
        sys.exit(f'源文件只有 {len(tbls)} 张表')
    tbl_el = copy.deepcopy(tbls[a.表 - 1]._tbl)

    if a.要:
        kws = [k.strip() for k in a.要.split(',') if k.strip()]
        rows = tbl_el.findall(TR)
        for tr in rows[1:]:
            tcs = tr.findall(TC)
            txt = cell_text(tcs[1]) if len(tcs) > 1 else ''
            if not any(k in txt for k in kws):
                tbl_el.remove(tr)

    dst_path = os.path.join(template_root(a.root), a.目标)
    os.makedirs(os.path.dirname(dst_path), exist_ok=True)
    if os.path.exists(dst_path):
        dst = Document(dst_path)
        dst.add_paragraph('')
    else:
        dst = Document()
    dst.add_paragraph(a.标题)

    copy_images(tbl_el, src.part, dst.part, {})
    anchor = dst.add_paragraph()
    anchor._p.addprevious(tbl_el)          # 必须插在 sectPr 之前，否则全堆到文末

    dst.save(dst_path)
    n_rows = len(tbl_el.findall(TR)) - 1
    n_img = len(tbl_el.findall('.//' + qn('w:drawing')))
    print(f'已入库：{a.目标}')
    print(f'  {a.标题}　{n_rows} 行，图 {n_img} 张')


if __name__ == '__main__':
    main()
